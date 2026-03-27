import pandas as pd
import openpyxl
import collections
import re
import random
import gc

wb_path = '研发领料明细-2025.1-12月-工厂做单（20260326）.xlsx'

prioritized_projects = [
    '8K双目全景追踪黑光摄像机', 'SMB系列AI专业安防摄像机', 'SMB系列澎湃算力智能NVR',
    '64路4盘位智能NVR', '8K智能摄像机', 'Wi-Fi双摄枪球联动摄像机',
    '4轴联动2400万枪球联动摄像机', '超低功耗门铃'
]

def classify_material(name, spec, project_names):
    name = str(name).lower() if name else ''
    spec = str(spec).lower() if spec else ''
    combined = name + ' ' + spec
    rules = [
        (r'nvr|录像机', lambda p: 'nvr' in p.lower()),
        (r'门铃', lambda p: '门铃' in p),
        (r'泛光灯', lambda p: '泛光灯' in p),
        (r'球|枪|三目', lambda p: '球' in p or '枪' in p),
        (r'太阳能|电池|低功耗|续航', lambda p: '太阳能' in p or '电池' in p or '低功耗' in p or '续航' in p),
        (r'8k|超清', lambda p: '8k' in p.lower() or '超清' in p),
        (r'全景', lambda p: '全景' in p),
        (r'动物', lambda p: '动物' in p),
        (r'存储|中心', lambda p: '存储' in p or '中心' in p)
    ]
    for pattern, filter_func in rules:
        if re.search(pattern, combined):
            matches = [p for p in project_names if filter_func(p)]
            if matches: return matches
    return project_names

wb = openpyxl.load_workbook(wb_path)
global_allocations = collections.defaultdict(lambda: collections.defaultdict(float))

for sheet_name in wb.sheetnames:
    print(f'Processing {sheet_name}...')
    ws = wb[sheet_name]

    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    if '数量校验' not in headers:
        continue

    start_col_idx = headers.index('数量校验') + 1
    project_names = [h for h in headers[start_col_idx:] if h]

    cycle_map = {}
    for i, p in enumerate(project_names):
        c_val = ws.cell(row=2, column=start_col_idx + 1 + i).value
        try:
            if c_val and '-' in str(c_val):
                start, end = str(c_val).split('-')
                cycle_map[p] = (int(start), int(end))
            else:
                cycle_map[p] = None
        except:
            cycle_map[p] = None

    doc_allocations = {}

    # Pre-read required data
    rows_data = []
    for r in range(4, ws.max_row + 1):
        date_val = ws.cell(row=r, column=1).value
        qty = ws.cell(row=r, column=13).value
        doc_no = ws.cell(row=r, column=4).value
        mat_code = ws.cell(row=r, column=9).value
        mat_name = ws.cell(row=r, column=10).value
        spec = ws.cell(row=r, column=11).value
        rows_data.append((r, date_val, qty, doc_no, mat_code, mat_name, spec))

    # Clear existing values
    for r in range(4, ws.max_row + 1):
        for c in range(start_col_idx + 1, ws.max_column + 1):
            ws.cell(row=r, column=c).value = None

    for r, date_val, qty, doc_no, mat_code, mat_name, spec in rows_data:
        date_str = str(date_val).strip()
        if not date_str or date_str == '合计' or date_str == 'None': continue

        try: qty = float(qty)
        except: continue

        item_ym = None
        try:
            if '/' in date_str:
                y, m, d = date_str.split('/')
                item_ym = int(f'{y}{int(m):02d}')
            elif '-' in date_str:
                y, m, d = date_str.split(' ')[0].split('-')
                item_ym = int(f'{y}{int(m):02d}')
        except: pass

        active_projects = project_names
        if item_ym:
            active_projects = []
            for p in project_names:
                c = cycle_map.get(p)
                if c is None: active_projects.append(p)
                elif c[0] <= item_ym <= c[1]: active_projects.append(p)
            if not active_projects:
                active_projects = project_names

        def pick(matches):
            valid = [m for m in matches if m in active_projects]
            if not valid: valid = matches
            prio = [p for p in valid if p in prioritized_projects]
            return random.choice(prio) if prio else random.choice(valid)

        doc_no = str(doc_no).strip()
        mat_code = str(mat_code).strip()
        mat_name = str(mat_name).strip()
        spec = str(spec).strip()

        if qty > 200000:
            n = len(active_projects)
            base = int(qty // n)
            rem = qty - base*n
            for i, p in enumerate(active_projects):
                q = base + rem if i == 0 else base
                c_idx = start_col_idx + 1 + project_names.index(p)
                ws.cell(row=r, column=c_idx).value = q
                global_allocations[(mat_code, mat_name)][p] += q
        elif qty > 0:
            selected = None
            if doc_no and doc_no != 'None':
                matches = classify_material(mat_name, spec, project_names)
                if doc_no in doc_allocations:
                    prev = doc_allocations[doc_no]
                    if prev in matches and prev in active_projects:
                        selected = prev
                    else:
                        selected = pick(matches)
                else:
                    selected = pick(matches)
                    doc_allocations[doc_no] = selected
            else:
                selected = pick(classify_material(mat_name, spec, project_names))
            ws.cell(row=r, column=start_col_idx + 1 + project_names.index(selected)).value = qty
            global_allocations[(mat_code, mat_name)][selected] += qty
        elif qty < 0:
            past = global_allocations[(mat_code, mat_name)]
            selected = None
            if past:
                best = max(past.items(), key=lambda x:x[1])[0]
                if best in project_names:
                    selected = best
            if not selected:
                selected = pick(classify_material(mat_name, spec, project_names))
            ws.cell(row=r, column=start_col_idx + 1 + project_names.index(selected)).value = qty
            global_allocations[(mat_code, mat_name)][selected] += qty

wb.save(wb_path)
print('Done!')
